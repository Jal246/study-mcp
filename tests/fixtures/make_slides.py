"""Builds notes/slides.pptx, the sample slide deck the tests read.

The .pptx is committed, so this only needs to run again if you change the deck:
    .venv/Scripts/python.exe tests/fixtures/make_slides.py
"""

from pathlib import Path

from pptx import Presentation
from pptx.util import Inches

OUT = Path(__file__).parent / "notes" / "slides.pptx"

deck = Presentation()
title_layout, bullets_layout, blank_layout = deck.slide_layouts[0], deck.slide_layouts[1], deck.slide_layouts[6]

# Slide 1: a title slide.
slide = deck.slides.add_slide(title_layout)
slide.shapes.title.text = "Graph Traversal"
slide.placeholders[1].text = "Week 7: exploring every vertex"

# Slide 2: bullet points, plus speaker notes (the lecturer's explanation, not shown on the slide).
slide = deck.slides.add_slide(bullets_layout)
slide.shapes.title.text = "Breadth-First Search"
body = slide.placeholders[1].text_frame
body.text = "Visit neighbours level by level"
body.add_paragraph().text = "Finds shortest paths in unweighted graphs"
slide.notes_slide.notes_text_frame.text = "BFS keeps the frontier in a queue: first in, first out."

# Slide 3: a heading inside a group shape, and a table. Naive text extraction misses both.
slide = deck.slides.add_slide(blank_layout)
group = slide.shapes.add_group_shape()
heading = group.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(8), Inches(0.8))
heading.text_frame.text = "BFS vs DFS"
table = slide.shapes.add_table(3, 2, Inches(0.5), Inches(1.5), Inches(8), Inches(2)).table
for row, cells in enumerate([("Algorithm", "Data structure"), ("BFS", "Queue"), ("DFS", "Stack")]):
    for col, text in enumerate(cells):
        table.cell(row, col).text = text

deck.save(OUT)
print(f"Wrote {OUT}")
