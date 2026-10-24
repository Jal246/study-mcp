"""Study Helper MCP server.

Gives an AI assistant (like Claude) tools to browse and search your study notes.
"""

import ctypes
import hashlib
import logging
import os
import re
import unicodedata
from contextlib import contextmanager
from dataclasses import dataclass
from datetime import date
from functools import lru_cache
from pathlib import Path

import numpy as np
from mcp.server.mcpserver import MCPServer
from mcp.server.mcpserver.exceptions import ToolError
import pypdfium2 as pdfium
import pypdfium2.raw as pdfium_c
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from study_mcp.cambria_math import SMALL_GLYPHS
from study_mcp.flashcards import Card, CardError, Deck
from study_mcp.quiz import QuizError, QuizLog

# Where the notes live. Override with the STUDY_NOTES_DIR environment variable.
DEFAULT_NOTES_DIR = Path(__file__).resolve().parents[2] / "notes"
NOTES_DIR = Path(os.environ.get("STUDY_NOTES_DIR", DEFAULT_NOTES_DIR)).resolve()

# File types we know how to read. Text files are one piece; PDFs and slide decks are split into
# numbered pieces, and this says what to call them: (file kind, piece name).
TEXT_EXTENSIONS = {".md", ".txt"}
PAGED_EXTENSIONS = {".pdf": ("PDF", "page"), ".pptx": ("PowerPoint", "slide")}
NOTE_EXTENSIONS = TEXT_EXTENSIONS | PAGED_EXTENSIONS.keys()

# read_note stops after this many characters so one huge PDF can't flood the AI's context.
MAX_READ_CHARS = 40_000

# Semantic search: a small embedding model that runs locally (no API key, works offline after
# the first ~70 MB download). It turns text into a list of 384 numbers (a "vector") that captures
# its meaning, so "how to find the remainder" lands close to a slide about "n mod d".
EMBED_MODEL = "BAAI/bge-small-en-v1.5"
# The downloaded model and saved vectors live here. Override with STUDY_CACHE_DIR.
DEFAULT_CACHE_DIR = Path(__file__).resolve().parents[2] / ".cache"
CACHE_DIR = Path(os.environ.get("STUDY_CACHE_DIR", DEFAULT_CACHE_DIR)).resolve()
MODEL_DIR = CACHE_DIR / "models"
VECTOR_DIR = CACHE_DIR / "vectors"
# Notes are split into chunks of about this many characters; each chunk gets its own vector.
CHUNK_CHARS = 1000
# Similarity scores go from -1 to 1, but with this model even unrelated text scores about 0.5,
# while real matches score about 0.65 and up. Results below this are dropped.
MIN_SCORE = 0.6

# Flashcards and their review history. Unlike .cache, this is the user's own work and can't be
# rebuilt, so it gets its own folder. Override with STUDY_DATA_DIR.
DEFAULT_DATA_DIR = Path(__file__).resolve().parents[2] / "data"
DATA_DIR = Path(os.environ.get("STUDY_DATA_DIR", DEFAULT_DATA_DIR)).resolve()

# The model download logs every HTTP request; keep that noise out of the server's logs.
logging.getLogger("httpx").setLevel(logging.WARNING)
os.environ.setdefault("HF_HUB_DISABLE_SYMLINKS_WARNING", "1")

mcp = MCPServer(
    name="study-helper",
    instructions=(
        "Tools for the user's personal study notes (text files, PDFs and PowerPoint slides). "
        "Use list_notes to see what exists. To find things, use semantic_search for questions and "
        "concepts in the user's own words (it matches meaning, not exact words), and search_notes "
        "for an exact term or formula. Use read_note to read a note (for PDFs you can read just "
        "some pages or slides). When answering from notes, say which note and page or slide the "
        "answer came from. "
        "Flashcards: when the user asks for cards, write short question/answer pairs from the "
        "notes and save each with add_flashcard (with its source). To quiz the user, call "
        "get_due_cards, ask ONE question at a time without showing the answer, wait for their "
        "reply, tell them if they were right, then grade_card it before the next question. "
        "Quizzes: when you write quiz questions yourself, save each graded answer with "
        "record_quiz_answer (with a short topic name). When the user asks what to study, or "
        "wants practice, check weak_topics first and focus on the weakest topics."
    ),
)


def _normalize(text: str) -> str:
    """Make text searchable.

    PDFs often store math as special Unicode letters: "𝑚𝑜𝑑" instead of "mod". They look the same
    but are different characters, so a search for "mod" would miss them. NFKC normalization turns
    those look-alikes back into plain letters.
    """
    return unicodedata.normalize("NFKC", text)


def _shape_lines(shapes) -> list[str]:
    """All text on a slide, in the order the shapes were added.

    Text hides in three places: text boxes (titles and bullets), tables, and groups of shapes,
    which can contain any of these, including more groups.
    """
    lines = []
    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            lines.extend(_shape_lines(shape.shapes))
        elif shape.has_text_frame:
            lines.append(shape.text_frame.text)
        elif shape.has_table:
            lines.extend(" | ".join(cell.text for cell in row.cells) for row in shape.table.rows)
    return lines


def _slide_texts(path: Path) -> list[str]:
    """One string per slide: the slide's text, then its speaker notes (often the lecturer's
    actual explanation of the slide)."""
    texts = []
    for slide in Presentation(str(path)).slides:
        lines = _shape_lines(slide.shapes)
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
            lines += ["Speaker notes:", slide.notes_slide.notes_text_frame.text]
        # PowerPoint stores a line break inside a paragraph (Shift+Enter) as a vertical tab.
        texts.append("\n".join(lines).replace("\v", "\n"))
    return texts


@dataclass
class _PdfChar:
    unit: int  # one UTF-16 code unit, as pdfium reports it
    font: str
    y: float  # baseline height on the page, to tell superscripts from subscripts
    height: float  # height of the glyph's box, to judge how far a raised character moved
    left: float
    right: float
    generated: bool  # a space or line break pdfium added, not text in the file


def _pdf_chars(textpage) -> list[_PdfChar]:
    chars = []
    for i in range(pdfium_c.FPDFText_CountChars(textpage.raw)):
        font = ctypes.create_string_buffer(256)
        pdfium_c.FPDFText_GetFontInfo(textpage.raw, i, font, len(font), None)
        x, y = ctypes.c_double(), ctypes.c_double()
        pdfium_c.FPDFText_GetCharOrigin(textpage.raw, i, ctypes.byref(x), ctypes.byref(y))
        box = [ctypes.c_double() for _ in range(4)]  # left, right, bottom, top
        pdfium_c.FPDFText_GetCharBox(textpage.raw, i, *(ctypes.byref(b) for b in box))
        chars.append(_PdfChar(
            pdfium_c.FPDFText_GetUnicode(textpage.raw, i),
            font.value.decode(errors="replace"),
            y.value,
            box[3].value - box[2].value,
            box[0].value,
            box[1].value,
            bool(pdfium_c.FPDFText_IsGenerated(textpage.raw, i)),
        ))
    return chars


def _shifted(char: _PdfChar, before: _PdfChar | None) -> bool:
    """Is `char` raised or lowered next to the character before it, like the 27 in 15²⁷ typed as
    ordinary text with superscript formatting? It must sit right after that character: a shift
    of more than a glyph's height, or a gap, means the text moved somewhere else on the page
    (like a slide number)."""
    if before is None or before.height <= 0:
        return False
    touching = abs(char.left - before.right) < 0.5 * before.height
    return touching and 0.1 * before.height < abs(char.y - before.y) < 0.7 * before.height


def _fix_math(chars: list[_PdfChar], math_fonts: set[str]) -> str:
    """A page's text, with superscripts and subscripts marked and broken equation glyphs repaired.

    PowerPoint and Word draw small equation glyphs (the 17 in 6¹⁷) from Cambria Math, and their
    PDF export writes the glyph's number instead of a character, so 6¹⁷ comes out as "6ଵ଻".
    We swap those back using SMALL_GLYPHS. Only fonts that also contain math letters (𝑥, 𝑚𝑜𝑑)
    are touched, so real Odia or Tamil text is left alone. Plain text has no superscripts, so
    without marks 6¹⁷ would read as 617; we mark them the way LaTeX does: "6^{17}",
    "p_{1}^{α_{1}}".
    """
    pieces: list[str] = []
    open_groups: list[str] = []  # "^" or "_" for each open "{"; index + 1 is the level
    base_y = [0.0, 0.0, 0.0]  # baseline of the last character at each level (0 = normal size)
    before: _PdfChar | None = None  # last normal-size character on this line

    def close_to(depth: int) -> None:
        while len(open_groups) > depth:
            open_groups.pop()
            pieces.append("}")

    for char in chars:
        text = chr(char.unit)
        if char.generated or text.isspace():
            close_to(0)
            pieces.append(text)
            if text in "\r\n":
                before = None
            continue
        small = SMALL_GLYPHS.get(char.unit) if char.font in math_fonts else None
        if small is not None:
            text, level = small
        else:
            # A normal glyph at the height of an open superscript is part of it: the ( in p^{min(α,β)}.
            level = next(
                (n for n in range(len(open_groups), 0, -1) if abs(char.y - base_y[n]) < 0.3), 0
            )
            if level == 0 and _shifted(char, before):
                level = 1
        if level == 0:
            close_to(0)
            pieces.append(text)
            base_y[0] = char.y
            before = char
            continue
        mark = "^" if char.y > base_y[level - 1] else "_"
        if len(open_groups) >= level and open_groups[level - 1] == mark:
            close_to(level)  # back from a deeper level: the "," after α₁ in min(α₁,β₁)
        else:
            close_to(level - 1)
            # pdfium often puts a line break between a letter and its raised exponent.
            while not open_groups and pieces and pieces[-1].isspace():
                pieces.pop()
            pieces.append(mark + "{")
            open_groups.append(mark)
        pieces.append(text)
        base_y[level] = char.y
    close_to(0)
    return "".join(pieces).encode("utf-16", "surrogatepass").decode("utf-16")


def _pdf_pages(path: Path) -> list[str]:
    with pdfium.PdfDocument(path) as pdf:
        pages = [_pdf_chars(page.get_textpage()) for page in pdf]
    # 0xD835 starts every math letter (𝑥, 𝑚, 𝛼): fonts using them are equation fonts.
    math_fonts = {char.font for page in pages for char in page if char.unit == 0xD835}
    return [_fix_math(page, math_fonts) for page in pages]


@lru_cache(maxsize=64)
def _load_pages(path: Path, modified_ns: int) -> list[str]:
    """Return a note's text as a list of pages or slides (text files count as one page).

    Text extraction is slow, so results are cached. `modified_ns` (the file's last-modified
    time) is part of the cache key: when the file changes, the key changes and we re-read it.
    """
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return [_normalize(text) for text in _pdf_pages(path)]
    if suffix == ".pptx":
        return [_normalize(text) for text in _slide_texts(path)]
    return [_normalize(path.read_text(encoding="utf-8"))]


def _pages(path: Path) -> list[str]:
    return _load_pages(path, path.stat().st_mtime_ns)


def _paged(path: Path) -> tuple[str, str] | None:
    """(file kind, piece name), e.g. ("PDF", "page"), for notes split into pieces; None for text."""
    return PAGED_EXTENSIONS.get(path.suffix.lower())


def _all_notes() -> list[Path]:
    """Every note file inside NOTES_DIR (including subfolders), sorted by path."""
    return sorted(
        p for p in NOTES_DIR.rglob("*") if p.is_file() and p.suffix.lower() in NOTE_EXTENSIONS
    )


def _name(path: Path) -> str:
    """A note's path relative to NOTES_DIR, with forward slashes, e.g. "math/module4.pdf"."""
    return path.relative_to(NOTES_DIR).as_posix()


def _where(path: Path, page_no: int, line_no: int) -> str:
    """Where a match is, for the user: "deck.pptx (slide 2)", "a.pdf (page 3)" or "notes.md:12"."""
    if paged := _paged(path):
        return f"{_name(path)} ({paged[1]} {page_no})"
    return f"{_name(path)}:{line_no}"


def _safe_path(relative_path: str) -> Path:
    """Turn a user-given path into a real path, refusing anything outside NOTES_DIR.

    Without this check, a request like "../../secret.txt" could read any file on the computer.
    """
    path = (NOTES_DIR / relative_path).resolve()
    if not path.is_relative_to(NOTES_DIR):
        raise ToolError(f"'{relative_path}' is outside the notes folder")
    if not path.is_file():
        raise ToolError(f"No note found at '{relative_path}'. Use list_notes to see available notes.")
    return path


@mcp.tool()
def list_notes() -> str:
    """List all study notes, with their paths and sizes (page or slide counts for PDFs and decks)."""
    notes = _all_notes()
    if not notes:
        return f"No notes found in {NOTES_DIR}"
    lines = []
    for p in notes:
        if paged := _paged(p):
            kind, piece = paged
            detail = f"{kind}, {len(_pages(p))} {piece}s"
        else:
            detail = f"{p.stat().st_size} bytes"
        lines.append(f"{_name(p)} ({detail})")
    return f"{len(notes)} notes in {NOTES_DIR}:\n" + "\n".join(lines)


@mcp.tool()
def read_note(path: str, start_page: int | None = None, end_page: int | None = None) -> str:
    """Read the text of one note. For PDFs and slide decks, optionally read only a range of
    pages (for decks, page numbers are slide numbers).

    Args:
        path: The note's path as shown by list_notes, e.g. "math/module4.pdf".
        start_page: First page or slide to read (starts at 1). Defaults to the first.
        end_page: Last page or slide to read, inclusive. Defaults to the last.
    """
    note = _safe_path(path)
    pages = _pages(note)
    if not (paged := _paged(note)):
        return pages[0]
    piece = paged[1]

    first = start_page or 1
    last = end_page or len(pages)
    if not 1 <= first <= last <= len(pages):
        raise ToolError(f"Invalid {piece} range {first}-{last}: '{path}' has {piece}s 1-{len(pages)}")

    parts = []
    total = 0
    for number in range(first, last + 1):
        chunk = f"--- {piece.title()} {number} ---\n{pages[number - 1].strip()}\n"
        if total + len(chunk) > MAX_READ_CHARS:
            parts.append(
                f"[Stopped at {piece} {number - 1} to stay under {MAX_READ_CHARS} characters. "
                f"Call read_note again with start_page={number} to continue.]"
            )
            break
        parts.append(chunk)
        total += len(chunk)
    return "\n".join(parts)


@mcp.tool()
def search_notes(query: str, max_results: int = 10) -> str:
    """Search all notes (including PDFs and slides) for a whole word or phrase, case-insensitive.

    Matches whole words only: "mod" finds "7 mod 3" but not "Module".

    Each result shows where the match is: "note:line" for text files, "note (page N)" for PDFs,
    "note (slide N)" for slide decks.

    Args:
        query: Text to look for, e.g. "chlorophyll".
        max_results: Maximum number of matching lines to return.
    """
    needle = _normalize(query).strip()
    if not needle:
        raise ToolError("query must not be empty")
    # (?<!\w) and (?!\w) mean "no letter/digit right before or after", i.e. whole words only.
    # re.escape makes characters like "(" or "+" in the query match literally.
    pattern = re.compile(rf"(?<!\w){re.escape(needle)}(?!\w)", re.IGNORECASE)

    matches = []
    for note in _all_notes():
        for page_no, page in enumerate(_pages(note), start=1):
            for line_no, line in enumerate(page.splitlines(), start=1):
                if not pattern.search(line):
                    continue
                matches.append(f"{_where(note, page_no, line_no)}: {line.strip()}")
                if len(matches) >= max_results:
                    return "\n".join(matches) + f"\n(stopped after {max_results} results)"

    return "\n".join(matches) if matches else f"No notes mention '{query}'."


@dataclass(frozen=True)
class Chunk:
    where: str  # Location to show the user, from _where: "note (page N)", "note:line", ...
    text: str


def _split(text: str, max_chars: int = CHUNK_CHARS) -> list[tuple[int, str]]:
    """Split text into pieces of about max_chars, cutting only between lines.

    Returns (line number the piece starts on, piece text) pairs, skipping blank pieces.
    """
    pieces = []
    current: list[str] = []
    size = 0
    start = 1
    for line_no, line in enumerate(text.splitlines(), start=1):
        if current and size + len(line) > max_chars:
            pieces.append((start, "\n".join(current)))
            current, size = [], 0
        if not current:
            start = line_no
        current.append(line)
        size += len(line) + 1
    if current:
        pieces.append((start, "\n".join(current)))
    return [(line_no, piece) for line_no, piece in pieces if piece.strip()]


def _chunks(path: Path) -> list[Chunk]:
    """Cut a note into chunks small enough to embed. Chunks never cross a page or slide boundary."""
    chunks = []
    for page_no, page in enumerate(_pages(path), start=1):
        for line_no, piece in _split(page):
            chunks.append(Chunk(_where(path, page_no, line_no), piece))
    return chunks


@lru_cache(maxsize=1)
def _embedder():
    """Load the embedding model once (it is downloaded on first use).

    Imported here, not at the top, so the server starts instantly and only pays for loading the
    model when semantic search is actually used.
    """
    from fastembed import TextEmbedding

    return TextEmbedding(EMBED_MODEL, cache_dir=str(MODEL_DIR))


def _unit(vectors: np.ndarray) -> np.ndarray:
    """Scale vectors to length 1, so a dot product between two of them is their cosine similarity."""
    return vectors / np.linalg.norm(vectors, axis=-1, keepdims=True)


@lru_cache(maxsize=64)
def _load_vectors(path: Path, modified_ns: int) -> tuple[list[Chunk], np.ndarray]:
    """Return a note's chunks and their vectors (a matrix with one row per chunk).

    Embedding is slow (seconds per PDF), so vectors are also saved to disk. The file name is a
    hash of the model name and the chunk texts: if the note or the model changes, the hash
    changes, and the old file is simply never looked up again.
    """
    chunks = _chunks(path)
    if not chunks:
        return [], np.empty((0, 0))
    fingerprint = "\n".join([EMBED_MODEL, *(c.where + "\n" + c.text for c in chunks)])
    saved = VECTOR_DIR / f"{hashlib.sha256(fingerprint.encode()).hexdigest()}.npy"
    if saved.exists():
        return chunks, np.load(saved)
    vectors = _unit(np.array(list(_embedder().embed([c.text for c in chunks]))))
    VECTOR_DIR.mkdir(parents=True, exist_ok=True)
    np.save(saved, vectors)
    return chunks, vectors


def _vectors(path: Path) -> tuple[list[Chunk], np.ndarray]:
    return _load_vectors(path, path.stat().st_mtime_ns)


@mcp.tool()
def semantic_search(query: str, max_results: int = 5) -> str:
    """Find the parts of the notes that are about what the query means, even if they use
    different words. E.g. "how to find the remainder" finds slides about "n mod d".

    Best for questions and concepts in your own words; for an exact word or formula, use
    search_notes. The first search can take a while (it downloads a small model and indexes
    the notes); later searches are fast.

    Each result shows a similarity score (higher = closer in meaning), where the passage is, and
    its start. Use read_note on that page to read the full context.

    Args:
        query: A question or description, e.g. "how do I check if a number is prime".
        max_results: Maximum number of passages to return.
    """
    question = _normalize(query).strip()
    if not question:
        raise ToolError("query must not be empty")
    # query_embed (not embed) adds the instruction prefix this model expects for search queries.
    query_vector = _unit(np.array(list(_embedder().query_embed([question])))[0])

    scored: list[tuple[float, Chunk]] = []
    for note in _all_notes():
        chunks, vectors = _vectors(note)
        if chunks:
            scored.extend(zip((vectors @ query_vector).tolist(), chunks))
    scored.sort(key=lambda pair: pair[0], reverse=True)

    results = []
    for score, chunk in scored[:max_results]:
        if score < MIN_SCORE:
            break
        snippet = " ".join(chunk.text.split())  # Collapse line breaks and runs of spaces.
        if len(snippet) > 300:
            snippet = snippet[:300] + "..."
        results.append(f"[{score:.2f}] {chunk.where}: {snippet}")
    return "\n\n".join(results) if results else f"Nothing in the notes seems related to '{query}'."


def _deck() -> Deck:
    return Deck(DATA_DIR / "flashcards.db")


@contextmanager
def _card_errors():
    """Turn flashcard problems (unknown id, bad grade...) into messages the AI can see."""
    try:
        yield
    except CardError as error:
        raise ToolError(str(error)) from error


def _next_review(card: Card) -> str:
    days = (card.due - date.today()).days
    when = "tomorrow" if days == 1 else f"in {days} days"
    return f"next review {when} ({card.due.isoformat()})"


@mcp.tool()
def add_flashcard(question: str, answer: str, source: str = "") -> str:
    """Save a flashcard for spaced-repetition review. New cards are due today.

    Keep cards small: one fact per card, a question with one clear answer.

    Args:
        question: E.g. "What data structure does BFS use?"
        answer: E.g. "A queue (first in, first out)."
        source: Where it came from, as shown by the search tools, e.g. "slides.pptx (slide 2)".
    """
    with _card_errors():
        card = _deck().add(question, answer, source, date.today())
    return f"Added card #{card.id}, due today."


@mcp.tool()
def get_due_cards(limit: int = 10) -> str:
    """List the flashcards due for review today, most overdue first, with their answers.

    The answers are for you to check the user's reply: ask the questions one at a time and don't
    reveal an answer until the user has tried. Then call grade_card.

    Args:
        limit: Maximum number of cards to return.
    """
    cards = _deck().due(date.today(), limit)
    if not cards:
        return "No cards due today."
    lines = [f"{len(cards)} cards due:"]
    for card in cards:
        source = f" [{card.source}]" if card.source else ""
        lines.append(f"#{card.id}{source}\nQ: {card.question}\nA: {card.answer}")
    return "\n\n".join(lines)


@mcp.tool()
def grade_card(card_id: int, quality: int) -> str:
    """Record how well the user answered a card; this schedules its next review.

    Args:
        card_id: The card's number, from get_due_cards.
        quality: 5 = instant and correct, 4 = correct after some thought, 3 = correct but hard,
            2 = wrong but the answer felt familiar, 1 = wrong, 0 = no idea.
    """
    with _card_errors():
        card = _deck().grade(card_id, quality, date.today())
    return f"Card #{card.id}: {_next_review(card)}."


@mcp.tool()
def delete_flashcard(card_id: int) -> str:
    """Delete a flashcard (e.g. one that is wrong or a duplicate) and its review history.

    Args:
        card_id: The card's number.
    """
    with _card_errors():
        card = _deck().delete(card_id)
    return f"Deleted card #{card.id}: {card.question}"


@mcp.tool()
def card_stats() -> str:
    """Summarize flashcard progress: totals, cards due, today's reviews, and the cards the user
    forgets most often (worth extra attention or rewriting)."""
    stats = _deck().stats(date.today())
    lines = [
        f"{stats['total']} cards: {stats['due']} due now, {stats['new']} never reviewed, "
        f"{stats['reviewed_today']} reviews done today."
    ]
    if stats["struggling"]:
        lines.append("Forgotten most often:")
        lines += [f"#{c.id} ({c.lapses}x forgotten): {c.question}" for c in stats["struggling"]]
    return "\n".join(lines)


def _quiz_log() -> QuizLog:
    return QuizLog(DATA_DIR / "quiz.db")


@mcp.tool()
def record_quiz_answer(
    topic: str, question: str, correct: bool, source: str = "", mistake: str = ""
) -> str:
    """Save how the user did on one quiz question, so weak_topics can find what to practice.

    Call this after grading each answer in a quiz you wrote (flashcards use grade_card instead).

    Args:
        topic: A short, reusable topic name, e.g. "RSA" or "Euclid's algorithm". Reuse the exact
            names weak_topics already shows, so answers on the same topic group together.
        question: The question asked.
        correct: Whether the user got it right.
        source: Where the question came from, e.g. "module4.pdf (page 25)".
        mistake: For wrong answers, what went wrong in a few words, e.g. "forgot the minus sign".
    """
    try:
        _quiz_log().record(topic, question, correct, source, mistake, date.today())
    except QuizError as error:
        raise ToolError(str(error)) from error
    return f"Saved: {'right' if correct else 'wrong'} on {topic.strip()}."


@mcp.tool()
def weak_topics(days: int | None = None) -> str:
    """Show the user's quiz scores by topic, weakest first, with their recent mistakes.

    Use it when the user asks what to study or practice, and before making a practice quiz:
    focus the questions on the weakest topics and their mistakes.

    Args:
        days: Only count the last this-many days (e.g. 7 for "this week"); omit for all time.
    """
    scores = _quiz_log().topics(date.today(), days)
    if not scores:
        return "No quiz answers saved yet."
    lines = []
    for s in scores:
        lines.append(
            f"{s.topic}: {s.correct}/{s.attempts} right ({s.accuracy:.0%}), "
            f"last quizzed {s.last_answered.isoformat()}"
        )
        lines += [f"  - missed: {m}" for m in s.recent_mistakes]
    return "\n".join(lines)
